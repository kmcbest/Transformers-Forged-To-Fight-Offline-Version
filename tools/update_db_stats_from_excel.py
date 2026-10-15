import openpyxl
import sqlite3
import math
import sys
import os

ROOT = r"d:\Agent\tftf"
EXCEL_PATH = os.path.join(ROOT, "TFTF_角色数值一览表.xlsx")
DB_PATH = os.path.join(ROOT, "Server", "tftf_database.db")

def round_to_significant_digits(d, digits=4):
    if d == 0:
        return 0
    num = math.floor(math.log10(abs(d))) + 1.0
    if num <= digits:
        return int(d)
    num2 = 10.0 ** num
    return int(num2 * round(float(d) / num2, digits))

def calculate_hero_rating(hp, atk, sig_lvl=100):
    base_hp = math.ceil(hp * 0.075)
    base_atk = math.ceil(atk * 1.0)
    r_hp = base_hp
    r_atk = base_atk
    if sig_lvl > 0:
        sig_bonus = 0.035 * math.log(max(1, sig_lvl))
        r_hp += math.ceil(base_hp * sig_bonus)
        r_atk += math.ceil(base_atk * sig_bonus)
    total_rating = round_to_significant_digits(r_hp + r_atk, 4)
    return total_rating

def update_db():
    print(f"Reading Excel: {EXCEL_PATH}")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    sheet = wb.active
    
    conn = sqlite3.connect(DB_PATH)
    c = conn.cursor()
    
    updates = []
    for r in range(2, sheet.max_row + 1):
        bid = sheet.cell(r, 2).value
        hp_j = sheet.cell(r, 10).value
        atk_k = sheet.cell(r, 11).value
        
        if not bid or hp_j is None or atk_k is None:
            continue
            
        hp = int(hp_j)
        atk = int(atk_k)
        rating = calculate_hero_rating(hp, atk, sig_lvl=100)
        updates.append((hp, atk, rating, bid))

    print(f"Prepared {len(updates)} updates for SQLite database...")
    c.executemany("""
        UPDATE characters
        SET hp = ?, attack = ?, rating = ?, updated_at = CURRENT_TIMESTAMP
        WHERE bot_id = ?
    """, updates)
    conn.commit()
    print("Database characters table successfully updated!")
    
    # Verify by reading back 5 samples
    print("\nVerification sample from database:")
    for row in c.execute("SELECT bot_id, name_zh, hp, attack, rating FROM characters LIMIT 5"):
        print(f"  {row[0]:<30} | {row[1]} | HP: {row[2]} | ATK: {row[3]} | Rating: {row[4]}")
        
    conn.close()

if __name__ == "__main__":
    update_db()
